import io
import json

import pytest

from app.file_extract import extract_text_from_bytes, is_supported_filename


def test_is_supported_filename():
    assert is_supported_filename("report.pdf")
    assert is_supported_filename("slides.PPTX")
    assert is_supported_filename("data.xlsx")
    assert not is_supported_filename("archive.zip")


def test_json_extract_pretty():
    raw = json.dumps({"a": 1, "b": [2, 3]}, ensure_ascii=False).encode("utf-8")
    text = extract_text_from_bytes(raw, "data.json")
    assert '"a": 1' in text
    assert '"b":' in text


def test_csv_extract_table():
    raw = "name,score\nAlice,10\nBob,20\n".encode("utf-8")
    text = extract_text_from_bytes(raw, "scores.csv")
    assert "Alice | 10" in text
    assert "Bob | 20" in text


def test_xlsx_extract():
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet["A1"] = "kpi"
    sheet["B1"] = 42
    buffer = io.BytesIO()
    workbook.save(buffer)
    text = extract_text_from_bytes(buffer.getvalue(), "metrics.xlsx")
    assert "Metrics" in text
    assert "kpi" in text
    assert "42" in text


def test_pptx_extract():
    pptx = pytest.importorskip("pptx")
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Hello slide"
    buffer = io.BytesIO()
    presentation.save(buffer)
    text = extract_text_from_bytes(buffer.getvalue(), "deck.pptx")
    assert "Слайд 1" in text
    assert "Hello slide" in text
