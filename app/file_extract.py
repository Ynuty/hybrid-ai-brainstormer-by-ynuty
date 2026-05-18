"""Extract plain text from uploaded context files for LLM prompts."""

from __future__ import annotations

import csv
import io
import json
import logging
from pathlib import PurePath

from app.config import get_settings

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "json",
    "csv",
    "py",
    "yaml",
    "yml",
    "xml",
    "html",
    "htm",
    "log",
}
BINARY_EXTENSIONS = {"pdf", "pptx", "ppt", "xlsx", "xlsm", "xls"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | BINARY_EXTENSIONS


def is_supported_filename(filename: str) -> bool:
    extension = _extension(filename)
    return extension in SUPPORTED_EXTENSIONS


def extract_text_from_bytes(raw: bytes, filename: str) -> str:
    extension = _extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        return f"[Файл не поддержан: .{extension or '?'} — допустимо: {supported}]"

    if extension == "pdf":
        return _pdf_to_text(raw, filename)
    if extension in {"pptx", "ppt"}:
        return _pptx_to_text(raw, filename)
    if extension in {"xlsx", "xlsm"}:
        return _xlsx_to_text(raw, filename)
    if extension == "xls":
        return _xls_to_text(raw, filename)
    if extension == "json":
        return _json_to_text(raw, filename)
    if extension == "csv":
        return _csv_to_text(raw, filename)
    return _text_to_string(raw, filename)


def _extension(filename: str) -> str:
    return PurePath(filename).suffix.lstrip(".").lower()


def _text_to_string(raw: bytes, filename: str) -> str:
    return raw.decode("utf-8", errors="replace").strip() or "_файл пуст_"


def _json_to_text(raw: bytes, filename: str) -> str:
    text = raw.decode("utf-8", errors="replace").strip()
    if not text:
        return "_JSON пуст_"
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        return f"[JSON не разобран: {exc}]\n\n{text[:20000]}"
    return json.dumps(data, ensure_ascii=False, indent=2)


def _csv_to_text(raw: bytes, filename: str) -> str:
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [
        " | ".join(cell.strip() for cell in row)
        for row in reader
        if any(cell.strip() for cell in row)
    ]
    if not rows:
        return "_CSV пуст_"
    return "### CSV\n\n" + "\n".join(rows[:5000])


def _pdf_to_text(raw: bytes, filename: str) -> str:
    try:
        import fitz
    except ImportError:
        return "[PDF не прочитан: установите pymupdf.]"

    try:
        document = fitz.open(stream=raw, filetype="pdf")
    except Exception as exc:
        logger.exception("Failed to open PDF %s: %s", filename, exc)
        return f"[PDF не прочитан: {exc}]"

    pages = []
    for page_index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if not text:
            text = _ocr_pdf_page(page, page_index)
        pages.append(f"### Страница {page_index}\n\n{text or '_текст не найден_'}")
    return "\n\n".join(pages) or "_PDF пуст_"


def _ocr_pdf_page(page, page_number: int) -> str:
    try:
        import fitz
        from PIL import Image
        import pytesseract
    except ImportError:
        return (
            f"[Страница {page_number}: OCR недоступен — установите pymupdf, pillow, pytesseract "
            "и системный Tesseract.]"
        )

    settings = get_settings()
    if settings.tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_path

    try:
        zoom = settings.pdf_ocr_dpi / 72
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        try:
            return pytesseract.image_to_string(image, lang="rus+eng").strip()
        except Exception:
            return pytesseract.image_to_string(image, lang="eng").strip()
    except Exception as exc:
        logger.exception("PDF OCR failed page=%s: %s", page_number, exc)
        return f"[Страница {page_number}: OCR не удалось: {exc}]"


def _pptx_to_text(raw: bytes, filename: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "[PPTX не прочитан: установите python-pptx.]"

    try:
        presentation = Presentation(io.BytesIO(raw))
    except Exception as exc:
        logger.exception("Failed to open PPTX %s: %s", filename, exc)
        return f"[PPTX не прочитан: {exc}]"

    slides: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            text = _shape_text(shape)
            if text:
                chunks.append(text)
        body = "\n\n".join(chunks) if chunks else "_слайд без текста_"
        slides.append(f"### Слайд {slide_index}\n\n{body}")
    return "\n\n".join(slides) or "_PPTX пуст_"


def _shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False) and shape.text_frame:
        paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs if p.text and p.text.strip()]
        if paragraphs:
            return "\n".join(paragraphs)
    if hasattr(shape, "text") and shape.text:
        return str(shape.text).strip()
    if getattr(shape, "has_table", False):
        try:
            table = shape.table
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                return "\n".join(rows)
        except Exception:
            return ""
    return ""


def _xlsx_to_text(raw: bytes, filename: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[Excel не прочитан: установите openpyxl.]"

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        logger.exception("Failed to open XLSX %s: %s", filename, exc)
        return f"[Excel не прочитан: {exc}]"

    sheets: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > 5000:
                rows.append("... [лист обрезан: больше 5000 строк] ...")
                break
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                rows.append(" | ".join(cells))
        body = "\n".join(rows) if rows else "_лист пуст_"
        sheets.append(f"### Лист: {sheet.title}\n\n{body}")
    workbook.close()
    return "\n\n".join(sheets) or "_Excel пуст_"


def _xls_to_text(raw: bytes, filename: str) -> str:
    try:
        import xlrd
    except ImportError:
        return "[XLS не прочитан: установите xlrd или сохраните файл как .xlsx.]"

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:
        logger.exception("Failed to open XLS %s: %s", filename, exc)
        return f"[XLS не прочитан: {exc}]"

    sheets: list[str] = []
    for sheet in book.sheets():
        rows: list[str] = []
        for row_index in range(min(sheet.nrows, 5000)):
            cells = [str(sheet.cell_value(row_index, col_index)).strip() for col_index in range(sheet.ncols)]
            if any(cells):
                rows.append(" | ".join(cells))
        body = "\n".join(rows) if rows else "_лист пуст_"
        sheets.append(f"### Лист: {sheet.name}\n\n{body}")
    return "\n\n".join(sheets) or "_Excel пуст_"
