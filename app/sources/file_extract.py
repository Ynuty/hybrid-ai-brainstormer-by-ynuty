"""Extract plain text from uploaded context files for LLM prompts."""

from __future__ import annotations

import csv
import io
import json
import logging
import subprocess
import tempfile
from pathlib import PurePath

from app.config import get_settings
from app.sources.ocr_utils import ocr_image_bytes, ocr_pil_image

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "json",
    "csv",
    "py",
    "yaml",
    "yml",
    "xml",
    "log",
}
HTML_EXTENSIONS = {"html", "htm"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff", "tif"}
DOCUMENT_EXTENSIONS = {"docx", "rtf", "odt", "epub", "doc"}
OFFICE_EXTENSIONS = {"pdf", "pptx", "ppt", "xlsx", "xlsm", "xls"}
AUDIO_EXTENSIONS = {"mp3", "wav", "m4a", "ogg", "flac", "webm", "mpeg", "mpga"}

BINARY_EXTENSIONS = OFFICE_EXTENSIONS | IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | BINARY_EXTENSIONS


def list_supported_file_extensions() -> list[str]:
    return sorted(SUPPORTED_EXTENSIONS | AUDIO_EXTENSIONS)


def is_audio_filename(filename: str) -> bool:
    return _extension(filename) in AUDIO_EXTENSIONS


def is_supported_filename(filename: str) -> bool:
    extension = _extension(filename)
    return extension in SUPPORTED_EXTENSIONS or extension in AUDIO_EXTENSIONS


def extract_text_from_bytes(raw: bytes, filename: str) -> str:
    extension = _extension(filename)
    if extension in AUDIO_EXTENSIONS:
        return "[Аудио обрабатывается отдельным транскрибатором на backend.]"

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS | AUDIO_EXTENSIONS))
        return f"[Файл не поддержан: .{extension or '?'} — допустимо: {supported}]"

    if extension == "pdf":
        return _pdf_to_text(raw, filename)
    if extension in {"pptx", "ppt"}:
        return _pptx_to_text(raw, filename)
    if extension in {"xlsx", "xlsm"}:
        return _xlsx_to_text(raw, filename)
    if extension == "xls":
        return _xls_to_text(raw, filename)
    if extension == "json":
        return _json_to_text(raw, filename)
    if extension == "csv":
        return _csv_to_text(raw, filename)
    if extension in HTML_EXTENSIONS:
        return _html_to_text(raw, filename)
    if extension == "docx":
        return _docx_to_text(raw, filename)
    if extension == "rtf":
        return _rtf_to_text(raw, filename)
    if extension == "odt":
        return _odt_to_text(raw, filename)
    if extension == "epub":
        return _epub_to_text(raw, filename)
    if extension == "doc":
        return _doc_to_text(raw, filename)
    if extension in IMAGE_EXTENSIONS:
        return _image_to_text(raw, filename)
    return _text_to_string(raw, filename)


def _extension(filename: str) -> str:
    return PurePath(filename).suffix.lstrip(".").lower()


def _text_to_string(raw: bytes, filename: str) -> str:
    return raw.decode("utf-8", errors="replace").strip() or "_файл пуст_"


def _json_to_text(raw: bytes, filename: str) -> str:
    text = raw.decode("utf-8", errors="replace").strip()
    if not text:
        return "_JSON пуст_"
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        return f"[JSON не разобран: {exc}]\n\n{text[:20000]}"
    return json.dumps(data, ensure_ascii=False, indent=2)


def _csv_to_text(raw: bytes, filename: str) -> str:
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [
        " | ".join(cell.strip() for cell in row)
        for row in reader
        if any(cell.strip() for cell in row)
    ]
    if not rows:
        return "_CSV пуст_"
    return "### CSV\n\n" + "\n".join(rows[:5000])


def _html_to_text(raw: bytes, filename: str) -> str:
    html = raw.decode("utf-8", errors="replace")
    if not html.strip():
        return "_HTML пуст_"
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return _text_to_string(raw, filename)

    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    return text or "_HTML без видимого текста_"


def _docx_to_text(raw: bytes, filename: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return "[DOCX не прочитан: установите python-docx.]"

    try:
        document = Document(io.BytesIO(raw))
    except Exception as exc:
        logger.exception("Failed to open DOCX %s: %s", filename, exc)
        return f"[DOCX не прочитан: {exc}]"

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    return "\n\n".join(paragraphs) or "_DOCX пуст_"


def _rtf_to_text(raw: bytes, filename: str) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError:
        return "[RTF не прочитан: установите striprtf.]"

    try:
        text = rtf_to_text(raw.decode("utf-8", errors="replace"))
    except Exception as exc:
        logger.exception("Failed to open RTF %s: %s", filename, exc)
        return f"[RTF не прочитан: {exc}]"
    return text.strip() or "_RTF пуст_"


def _odt_to_text(raw: bytes, filename: str) -> str:
    try:
        from odf.opendocument import load
        from odf import text as odf_text
        from odf.teletype import extractText
    except ImportError:
        return "[ODT не прочитан: установите odfpy.]"

    try:
        with tempfile.NamedTemporaryFile(suffix=".odt", delete=False) as tmp:
            tmp.write(raw)
            tmp.flush()
            document = load(tmp.name)
    except Exception as exc:
        logger.exception("Failed to open ODT %s: %s", filename, exc)
        return f"[ODT не прочитан: {exc}]"

    chunks: list[str] = []
    for element in document.getElementsByType(odf_text.P):
        text = extractText(element)
        if text and text.strip():
            chunks.append(text.strip())
    return "\n\n".join(chunks) or "_ODT пуст_"


def _epub_to_text(raw: bytes, filename: str) -> str:
    try:
        from ebooklib import ITEM_DOCUMENT, epub
    except ImportError:
        return "[EPUB не прочитан: установите ebooklib.]"

    try:
        book = epub.read_epub(io.BytesIO(raw))
    except Exception as exc:
        logger.exception("Failed to open EPUB %s: %s", filename, exc)
        return f"[EPUB не прочитан: {exc}]"

    try:
        from bs4 import BeautifulSoup
    except ImportError:
        BeautifulSoup = None

    parts: list[str] = []
    for item in book.get_items():
        if item.get_type() != ITEM_DOCUMENT:
            continue
        content = item.get_content().decode("utf-8", errors="replace")
        if BeautifulSoup:
            soup = BeautifulSoup(content, "lxml")
            text = soup.get_text("\n", strip=True)
        else:
            text = content
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts) or "_EPUB пуст_"


def _doc_to_text(raw: bytes, filename: str) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
            tmp.write(raw)
            tmp.flush()
            result = subprocess.run(
                ["antiword", tmp.name],
                capture_output=True,
                text=True,
                timeout=30,
                check=False,
            )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()
    except FileNotFoundError:
        pass
    except Exception as exc:
        logger.warning("antiword failed for %s: %s", filename, exc)
    return (
        "[Формат .doc не поддержан на этом сервере. "
        "Сохраните документ как .docx или экспортируйте в PDF.]"
    )


def _image_to_text(raw: bytes, filename: str) -> str:
    return f"### Изображение: {filename}\n\n{ocr_image_bytes(raw)}"


def _pdf_to_text(raw: bytes, filename: str) -> str:
    try:
        import fitz
    except ImportError:
        return "[PDF не прочитан: установите pymupdf.]"

    try:
        document = fitz.open(stream=raw, filetype="pdf")
    except Exception as exc:
        logger.exception("Failed to open PDF %s: %s", filename, exc)
        return f"[PDF не прочитан: {exc}]"

    pages = []
    for page_index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if not text:
            text = _ocr_pdf_page(page, page_index)
        pages.append(f"### Страница {page_index}\n\n{text or '_текст не найден_'}")
    return "\n\n".join(pages) or "_PDF пуст_"


def _ocr_pdf_page(page, page_number: int) -> str:
    try:
        import fitz
        from PIL import Image
    except ImportError:
        return (
            f"[Страница {page_number}: OCR недоступен — установите pymupdf, pillow, pytesseract "
            "и системный Tesseract.]"
        )

    settings = get_settings()
    try:
        zoom = settings.pdf_ocr_dpi / 72
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        return ocr_pil_image(image) or f"_Страница {page_number}: текст не найден_"
    except Exception as exc:
        logger.exception("PDF OCR failed page=%s: %s", page_number, exc)
        return f"[Страница {page_number}: OCR не удалось: {exc}]"


def _pptx_to_text(raw: bytes, filename: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "[PPTX не прочитан: установите python-pptx.]"

    try:
        presentation = Presentation(io.BytesIO(raw))
    except Exception as exc:
        logger.exception("Failed to open PPTX %s: %s", filename, exc)
        return f"[PPTX не прочитан: {exc}]"

    slides: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            text = _shape_text(shape)
            if text:
                chunks.append(text)
        body = "\n\n".join(chunks) if chunks else "_слайд без текста_"
        slides.append(f"### Слайд {slide_index}\n\n{body}")
    return "\n\n".join(slides) or "_PPTX пуст_"


def _shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False) and shape.text_frame:
        paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs if p.text and p.text.strip()]
        if paragraphs:
            return "\n".join(paragraphs)
    if hasattr(shape, "text") and shape.text:
        return str(shape.text).strip()
    if getattr(shape, "has_table", False):
        try:
            table = shape.table
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                return "\n".join(rows)
        except Exception:
            return ""
    return ""


def _xlsx_to_text(raw: bytes, filename: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[Excel не прочитан: установите openpyxl.]"

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        logger.exception("Failed to open XLSX %s: %s", filename, exc)
        return f"[Excel не прочитан: {exc}]"

    sheets: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > 5000:
                rows.append("... [лист обрезан: больше 5000 строк] ...")
                break
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                rows.append(" | ".join(cells))
        body = "\n".join(rows) if rows else "_лист пуст_"
        sheets.append(f"### Лист: {sheet.title}\n\n{body}")
    workbook.close()
    return "\n\n".join(sheets) or "_Excel пуст_"


def _xls_to_text(raw: bytes, filename: str) -> str:
    try:
        import xlrd
    except ImportError:
        return "[XLS не прочитан: установите xlrd или сохраните файл как .xlsx.]"

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:
        logger.exception("Failed to open XLS %s: %s", filename, exc)
        return f"[XLS не прочитан: {exc}]"

    sheets: list[str] = []
    for sheet in book.sheets():
        rows: list[str] = []
        for row_index in range(min(sheet.nrows, 5000)):
            cells = [str(sheet.cell_value(row_index, col_index)).strip() for col_index in range(sheet.ncols)]
            if any(cells):
                rows.append(" | ".join(cells))
        body = "\n".join(rows) if rows else "_лист пуст_"
        sheets.append(f"### Лист: {sheet.name}\n\n{body}")
    return "\n\n".join(sheets) or "_Excel пуст_"
